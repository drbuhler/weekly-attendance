#!/usr/bin/env python3
"""Check that the in-browser parser (roll.js, run under Node) and the Python export+merge
(tools/export_weekly_tabs.py + tools/build_encrypted.py) agree EXACTLY on a workbook.

    python3 tools/parity_check.py path/to/attendance.xlsx [--now 2026-10-11T11:59:00-07:00]

The workbook's tabs are converted to Google-style CSV in a temp dir, roll.js builds the roster
from them, and both sides are compared by SHA-256 digest. Prints aggregate counts only.
Needs: node, openpyxl, cryptography.
"""
import argparse, csv, datetime as dt, hashlib, json, os, subprocess, sys, tempfile, shutil

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_encrypted  # noqa: E402

NODE_SNIPPET = r"""
const fs = require('fs'), path = require('path'), crypto = require('crypto');
const R = require(process.argv[1]); const dir = process.argv[2]; const now = new Date(process.argv[3]);
const tabs = JSON.parse(fs.readFileSync(path.join(dir, 'tabs.json'), 'utf8'));
const html = tabs.map(t => `items.push({name: ${JSON.stringify(t.name)}, pageUrl: "x", gid: "${t.gid}",initialSheet: false});`).join('');
const found = R.discoverTabs(html);
const sel = R.selectWeeklyTabs(found, now, 8);
const weekly = sel.map(t => ({ name: t.name, date: t.date, rows: R.parseCSV(fs.readFileSync(path.join(dir, t.gid + '.csv'), 'utf8')) }));
const out = R.build(weekly);
const digest = crypto.createHash('sha256').update(JSON.stringify(out.people.map(p => [p.n, p.w]))).digest('hex');
console.log(JSON.stringify({ tabs: sel.map(t => t.name), weeks: out.weeks, n: out.people.length, digest, unknownMarks: out.unknownMarks,
  codes: out.people.map(p => p.w) }));
"""


def cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def status_counts(codes):
    red = yellow = 0
    for w in codes:
        s = 0
        for c in reversed(w):
            if c == "A": s += 1
            elif c == "G": continue
            else: break
        red += s >= 5; yellow += 3 <= s <= 4
    return red, yellow, len(codes) - red - yellow


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--now", default=dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                    help="instant to evaluate the noon-Pacific tab rule at (ISO with offset)")
    a = ap.parse_args()
    tmp = tempfile.mkdtemp(prefix="parity-")
    os.chmod(tmp, 0o700)
    try:
        wb = openpyxl.load_workbook(a.xlsx, data_only=True, read_only=True)
        tabs = []
        for gid, name in enumerate(wb.sheetnames):
            tabs.append({"name": name, "gid": str(gid)})
            with open(os.path.join(tmp, f"{gid}.csv"), "w", newline="") as f:
                w = csv.writer(f, lineterminator="\r\n")
                for row in wb[name].iter_rows(values_only=True):
                    w.writerow([cell(v) for v in row])
        json.dump(tabs, open(os.path.join(tmp, "tabs.json"), "w"))
        js = json.loads(subprocess.run(["node", "-e", NODE_SNIPPET, os.path.join(ROOT, "roll.js"), tmp, a.now],
                                       check=True, capture_output=True, text=True).stdout)
        # Python side
        out_csv = os.path.join(tmp, "py.csv")
        subprocess.run([sys.executable, os.path.join(HERE, "export_weekly_tabs.py"), a.xlsx, out_csv, "--now", a.now],
                       check=True, capture_output=True)
        with open(out_csv, newline="") as f:
            first_tab = next(csv.reader(f))[1]
        year = js["weeks"][0]["date"][:4]  # same start year the export inferred
        weeks, people, _ = build_encrypted.load(out_csv, int(year))
        build_encrypted.load_visitors(os.path.join(tmp, "py.visitors.csv"), weeks, int(year))
        pd = hashlib.sha256(json.dumps([[p["n"], p["w"]] for p in people], ensure_ascii=False,
                                       separators=(",", ":")).encode()).hexdigest()
        py_vis = [w["visitors"] for w in weeks]
        js_vis = [w["visitors"] for w in js["weeks"]]
        same_weeks = [(w["label"], w["date"]) for w in weeks] == [(w["label"], w["date"]) for w in js["weeks"]]
        r, y, o = status_counts(js["codes"])
        last = len(weeks) - 1
        pres = sum(c[last] == "P" for c in js["codes"]); onsheet = sum(c[last] in "PA" for c in js["codes"])
        print(json.dumps({"tabs_used": js["tabs"], "people_js": js["n"], "people_py": len(people),
                          "roster_digest_match": pd == js["digest"], "weeks_match": same_weeks,
                          "visitors_js": js_vis, "visitors_py": py_vis, "red_yellow_ok": [r, y, o],
                          "latest_attended": f"{pres}/{onsheet}", "unknown_marks_js": js["unknownMarks"]}))
        if pd != js["digest"] or py_vis != js_vis or not same_weeks:
            sys.exit(1)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
