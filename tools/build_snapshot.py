#!/usr/bin/env python3
"""Build data.enc.json: the encrypted fallback snapshot PLUS the encrypted live-sheet settings.

The page shows this snapshot instantly after the passcode, then switches to the live sheet if it
loads. If the live sheet can't be read (Google outage, unpublished sheet, network, layout change),
the snapshot stays on screen with a "showing data as of ..." note. Re-run this whenever you
change the passcode or the sheet link, and now and then to keep the fallback fresh:

    python3 tools/build_snapshot.py            # downloads the published sheet (.xlsx)
    python3 tools/build_snapshot.py --xlsx FILE  # or use a local .xlsx download

Reads ATTENDANCE_SHEET_URL (the "Publish to the web" link) and ATTENDANCE_PASSCODE from the
environment, or prompts for them with hidden input. Neither is printed or stored in plain text.
Roster files only ever exist in a private temp dir that is deleted afterwards. Prints counts only.
"""
import argparse, csv, datetime as dt, getpass, json, os, re, shutil, subprocess, sys, tempfile, urllib.request

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_encrypted as be  # noqa: E402
from names import name_key  # noqa: E402
import older_record  # noqa: E402

PUB_RE = re.compile(r"^https://docs\.google\.com/spreadsheets/d/e/(2PACX-[A-Za-z0-9_-]{20,})(?:/[^?#]*)?(?:[?#].*)?$")


def archive_rows(xlsx):
    """The 'Archive' tab (written by tools/weekly_tab.gs) as text rows, like its published CSV."""
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "archive"), None)
    if ws is None:
        return []
    def txt(v):
        if v is None:
            return ""
        if isinstance(v, bool):
            return "TRUE" if v else "FALSE"
        if isinstance(v, (dt.datetime, dt.date)):
            return v.strftime("%Y-%m-%d")
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    rows = [[txt(v) for v in r] for r in ws.iter_rows(values_only=True)]
    return [r for r in rows if any(c.strip() for c in r)]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", help="use this local .xlsx instead of downloading the published sheet")
    ap.add_argument("-o", "--out", default=os.path.join(ROOT, "data.enc.json"))
    ap.add_argument("--title", default="Catechumen attendance")
    ap.add_argument("--weeks", type=int, default=8)
    ap.add_argument("--now", help="pretend the current time is this ISO date-time (testing)")
    ap.add_argument("--iterations", type=int, default=be.DEFAULT_ITER)
    args = ap.parse_args()

    url = os.environ.get("ATTENDANCE_SHEET_URL")
    if url is None:
        if not sys.stdin.isatty():
            sys.exit("Set ATTENDANCE_SHEET_URL or run interactively to be prompted.")
        url = getpass.getpass("Published sheet URL (hidden): ")
    m = PUB_RE.match(url.strip())
    if not m:
        sys.exit("That doesn't look like a 'Publish to the web' Google Sheets link (…/spreadsheets/d/e/2PACX-…/pubhtml).")
    base = f"https://docs.google.com/spreadsheets/d/e/{m[1]}/"
    del url
    passcode = be.get_passcode()

    tmp = tempfile.mkdtemp(prefix="snapshot-")
    os.chmod(tmp, 0o700)
    try:
        xlsx = args.xlsx
        if not xlsx:
            xlsx = os.path.join(tmp, "sheet.xlsx")
            with urllib.request.urlopen(base + "pub?output=xlsx", timeout=60) as r, open(xlsx, "wb") as f:
                shutil.copyfileobj(r, f)
        csv_path = os.path.join(tmp, "roll.csv")
        cmd = [sys.executable, os.path.join(HERE, "export_weekly_tabs.py"), xlsx, csv_path, "--weeks", str(args.weeks)]
        if args.now:
            cmd += ["--now", args.now]
        r = subprocess.run(cmd, capture_output=True, text=True)
        if r.returncode:
            sys.exit("Export failed.")
        with open(csv_path, newline="", encoding="utf-8") as f:
            first = f.readline().split(",")[1].strip()
        today = dt.date.today()
        year = today.year - (1 if be.MONTHS.get(first[:3].lower(), 0) > today.month else 0)
        weeks, people, _ = be.load(csv_path, year)
        be.load_visitors(os.path.join(tmp, "roll.visitors.csv"), weeks, year)
        with open(os.path.join(tmp, "roll.baptized.csv"), newline="", encoding="utf-8") as f:
            marked = {name_key(r[0]) for r in list(csv.reader(f))[1:] if r}
        for p in people:
            if name_key(p["n"]) in marked:
                p["b"] = 1
        archive = archive_rows(xlsx)
        # older attendance (master grids, summer tabs, weekly tabs) for "Last seen <Mon YYYY>"
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
        roster = [p["n"] for p in people if p["w"][-1] in "PA"]
        older, ostats = older_record.build(wb, [p["n"] for p in people], dt.date.fromisoformat(weeks[-1]["date"]), roster)
        ostats.pop("rules", None)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    payload = {"v": 1, "title": args.title, "banner": "",
               "generated": dt.datetime.now().astimezone().isoformat(timespec="minutes"),
               "weeks": weeks, "people": people, "archive": archive, "older": older,
               "live": {"pub": base, "weeks": args.weeks}}
    blob = be.encrypt(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
                      passcode, args.iterations)
    del passcode, payload, base
    out = os.path.abspath(args.out)
    fd, tmpf = tempfile.mkstemp(dir=os.path.dirname(out), prefix=".enc-", suffix=".tmp")
    with os.fdopen(fd, "w") as f:
        json.dump(blob, f)
        f.write("\n")
    os.chmod(tmpf, 0o644)
    os.replace(tmpf, out)
    print(f"Wrote {out}", file=sys.stderr)
    print(be.summarize(weeks, people), file=sys.stderr)
    print("older record: " + ", ".join(f"{k} {v}" for k, v in ostats.items()), file=sys.stderr)


if __name__ == "__main__":
    main()
